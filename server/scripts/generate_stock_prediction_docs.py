import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Page setup - Standard Letter with 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Palette
    NAVY_HEX = "0B1736"
    BLUE_HEX = "1D4ED8"
    LIGHT_BG = "F3F4F6"
    BORDER_HEX = "D1D5DB"
    NAVY_COLOR = RGBColor(11, 23, 54)
    BLUE_COLOR = RGBColor(29, 78, 216)
    DARK_COLOR = RGBColor(31, 41, 55)
    GRAY_COLOR = RGBColor(107, 114, 128)

    # --- Title Block ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("InvenSight — Stock Prediction & Inventory Forecasting Architecture")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY_COLOR

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Technical Documentation, Algorithmic Engine & Operational Guide for JonBrix Motor Parts")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = GRAY_COLOR

    # Meta Table
    meta_table = doc.add_table(rows=2, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [("System Version", "InvenSight v1.0 Production"), ("Target Entity", "JonBrix Motor Parts"), ("Module", "Stock Prediction & Forecasting"), ("Date", "August 2026")],
        [("Core Engine", "services/analytics_engine.py"), ("Algorithm Stack", "Meta Prophet + Multi-Timeframe Velocity"), ("Classification", "Internal Technical Reference"), ("Status", "Active Production")]
    ]
    for row_idx, row_data in enumerate(meta_data):
        for col_idx, (k, v) in enumerate(row_data):
            cell = meta_table.rows[row_idx].cells[col_idx]
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r1 = p.add_run(f"{k}\n")
            r1.font.size = Pt(8.5)
            r1.font.bold = True
            r1.font.color.rgb = GRAY_COLOR
            r2 = p.add_run(v)
            r2.font.size = Pt(9.5)
            r2.font.bold = True
            r2.font.color.rgb = NAVY_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = NAVY_COLOR
        return p

    def add_sub_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = BLUE_COLOR
        return p

    def add_body(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Arial"
            rb.font.size = Pt(10)
            rb.font.bold = True
            rb.font.color.rgb = DARK_COLOR
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.italic = italic
        r.font.color.rgb = DARK_COLOR
        return p

    def add_callout(text, title="KEY PRINCIPLE"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        set_cell_background(cell, "EFF6FF")
        set_cell_margins(cell, 140, 140, 180, 180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        rt = p.add_run(f"📌 {title}: ")
        rt.font.name = "Arial"
        rt.font.size = Pt(9.5)
        rt.font.bold = True
        rt.font.color.rgb = BLUE_COLOR
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = DARK_COLOR
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- Section 1: Executive Overview ---
    add_section_header("1. Executive Overview & Purpose")
    add_body("The InvenSight Stock Prediction Engine is an automated decision-support system designed specifically for retail motorcycle parts, lubricants, and aftermarket accessories at JonBrix. In traditional retail environments, inventory managers face two major financial threats:")
    add_body("Stockouts happen when popular consumables (such as engine oils, brake pads, and drive belts) run out. Every day an item is out of stock results in immediate lost revenue and customer defection to competitors.", "1. Lost Sales (Under-Stocking): ")
    add_body("Ordering excessive inventory ties up cashflow, consumes limited warehouse shelf space, and exposes products to damage, aging, or obsolete capital.", "2. Tied Capital & Deadstock (Over-Stocking): ")
    add_body("To solve this, InvenSight continuously analyzes real-time Point of Sale (POS) transactions, historical velocity, supplier lead times, and demand momentum to predict exactly when stock will deplete and how many units should be reordered.")

    add_callout("The Stock Prediction Engine operates completely autonomously on every POS transaction, recalculating sales velocity, days until stockout, and purchase order quantities in real-time.", "AUTOMATED ENGINE")

    # --- Section 2: Data Ingestion & Sales Velocity ---
    add_section_header("2. Multi-Timeframe Sales Velocity Model")
    add_body("Rather than relying on a single static average, InvenSight ingests sales data across four distinct analytical windows to balance immediate short-term demand surges against long-term seasonality:")

    # Velocity Table
    tbl_vel = doc.add_table(rows=5, cols=3)
    tbl_vel.alignment = WD_TABLE_ALIGNMENT.CENTER
    vel_headers = ["Timeframe Metric", "Variable Name", "Operational Significance"]
    for col_idx, text in enumerate(vel_headers):
        c = tbl_vel.rows[0].cells[col_idx]
        set_cell_background(c, "0B1736")
        set_cell_margins(c, 100, 100, 120, 120)
        p = c.paragraphs[0]
        r = p.add_run(text)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    vel_data = [
        ("Recent 14-Day Sales", "sold14", "Captures current high-frequency consumer demand over the past 2 weeks (daily_14 = sold14 / 14.0)."),
        ("Previous 14-Day Sales", "sold_prev14", "Used as a benchmark against sold14 to calculate the 2-week Demand Momentum Growth Ratio."),
        ("Monthly 30-Day Sales", "sold30", "Provides monthly baseline velocity for general demand smoothing (daily_30 = sold30 / 30.0)."),
        ("Annual 365-Day Sales", "sold365", "Accounts for 1-year historical volume and seasonality across peak holiday or rainy seasons.")
    ]
    for row_idx, data in enumerate(vel_data, start=1):
        for col_idx, text in enumerate(data):
            c = tbl_vel.rows[row_idx].cells[col_idx]
            set_cell_background(c, "FFFFFF" if row_idx % 2 != 0 else "F9FAFB")
            set_cell_margins(c, 80, 80, 100, 100)
            p = c.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9)
            r.font.color.rgb = DARK_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_sub_header("Out-of-Stock Demand Fallback (Preventing False Zeros)")
    add_body("When a product runs out of stock, standard Point of Sale systems record 0 sales simply because there is nothing on the shelves to buy. A naive algorithm would mistakenly assume demand has dropped to zero. InvenSight solves this by automatically switching to a weighted historical blend when current_stock <= 0:")
    add_body("Daily Demand (Out of Stock) = (daily_30 × 25%) + (daily_365 × 25%) + (daily_all_time × 50%)", "Formula: ", italic=True)
    add_body("This ensures that out-of-stock bestsellers immediately trigger high-priority reorder suggestions with accurate quantities.")

    # --- Section 3: Machine Learning ---
    add_section_header("3. Machine Learning Layer — Meta Prophet Time-Series Model")
    add_body("For top high-volume revenue generators (Top 5 SKUs with >= 90 days of recorded sales history), InvenSight deploys an automated Meta Prophet forecasting model. Prophet decomposes historical time-series data into:")
    add_body("Identifies long-term growth or decline over a 730-day (2-year) training window.", "• Trend Component: ")
    add_body("Models recurring weekly spikes (e.g., higher mechanic repairs and oil changes on Saturdays and Sundays).", "• Day-of-Week Seasonality: ")
    add_body("Detects payday surges (15th and 30th of the month) and holiday spikes.", "• Calendar Seasonality: ")
    add_body("Computes an 80% confidence interval, allowing the engine to calculate a mathematical certainty score for every recommendation.", "• Uncertainty Estimation: ")

    add_callout("Safety Floor Rule: If Prophet outputs an unusually low forecast due to temporary stock disruptions, the engine automatically applies a floor: daily = max(prophet_forecast, daily_velocity). The system never under-predicts demand.", "ALGORITHMIC SAFETY GUARD")

    # --- Section 4: Demand Momentum Factor ---
    add_section_header("4. Two-Week Demand Momentum Factor")
    add_body("To adapt dynamically to sudden market shifts (such as viral product trends, supplier promotions, or weather changes), InvenSight calculates a Demand Momentum Factor by comparing sales over the latest 14 days against the preceding 14 days:")
    add_body("Growth Ratio (R) = sold14 / sold_prev14", "Formula: ", italic=True)

    # Momentum Table
    tbl_mom = doc.add_table(rows=5, cols=4)
    tbl_mom.alignment = WD_TABLE_ALIGNMENT.CENTER
    mom_headers = ["Market Trend State", "Growth Ratio (R)", "Momentum Multiplier", "Inventory Buffer Action"]
    for col_idx, text in enumerate(mom_headers):
        c = tbl_mom.rows[0].cells[col_idx]
        set_cell_background(c, "0B1736")
        set_cell_margins(c, 100, 100, 120, 120)
        p = c.paragraphs[0]
        r = p.add_run(text)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    mom_data = [
        ("Surging / Rapid Growth", "R >= 1.25 (+25% growth)", "1.25x (125%)", "Adds 25% extra safety stock buffer to prevent stockouts."),
        ("Growing / High Demand", "1.00 <= R < 1.25", "1.15x (115%)", "Adds 15% safety buffer for steady, reliable inventory turnover."),
        ("Normal Steady Pace", "0.70 <= R < 1.00", "1.05x (105%)", "Applies standard 5% buffer for regular replenishment."),
        ("Cooling / Dropping Demand", "R < 0.70 (>=30% drop)", "0.85x (85%)", "Downscales reorder recommendation by 15% to prevent overstocking.")
    ]
    for row_idx, data in enumerate(mom_data, start=1):
        for col_idx, text in enumerate(data):
            c = tbl_mom.rows[row_idx].cells[col_idx]
            set_cell_background(c, "FFFFFF" if row_idx % 2 != 0 else "F9FAFB")
            set_cell_margins(c, 80, 80, 100, 100)
            p = c.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9)
            r.font.color.rgb = DARK_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- Section 5: Days to Stockout & Reorder Formulas ---
    add_section_header("5. Days to Stockout & Automated Reorder Formulas")
    add_sub_header("A. Days Until Stockout")
    add_body("Days to Stockout calculates the exact number of days remaining before on-hand inventory hits zero:")
    add_body("Days to Stockout = Current Stock / (Daily Sales Velocity × Momentum Factor)", "Formula: ", italic=True)

    add_sub_header("B. Target 14-Day (2-Week) Stock Level")
    add_body("To maintain optimal buffer without tying up excess capital, the target stocking level covers 14 days of projected demand plus safety buffer:")
    add_body("Target Stock (14 Days) = (14-Day Projected Demand) + max(Reorder Level, 3 Days Velocity)", "Formula: ", italic=True)

    add_sub_header("C. Recommended Reorder Quantity")
    add_body("The algorithm dynamically determines the exact number of units to order based on actual units sold demand:")
    add_body("Recommended Reorder Qty = max(ceil(Target Stock), ceil(sold14 × 1.15), Reorder Level × 2)", "• For Out of Stock Items: ", italic=True)
    add_body("If Current Stock < Target Stock: Reorder Qty = ceil(Target Stock - Current Stock). If Current Stock >= Target Stock: Reorder Qty = 0 (No purchase needed).", "• For In-Stock Items: ", italic=True)

    # --- Section 6: Risk Classification ---
    add_section_header("6. Risk Levels & Multi-Period Horizon Simulation")
    add_body("Every inventory item in JonBrix is classified into one of three visual risk tiers:")
    add_body("Days to Stockout <= 7 days, or Current Stock = 0. Requires immediate purchase order placement to supplier.", "🔴 Critical / High Risk: ")
    add_body("Days to Stockout between 8 and 14 days, or Current Stock <= Reorder Level. Needs review in the upcoming weekly supplier batch.", "🟡 Medium Risk / Warning: ")
    add_body("Days to Stockout >= 15 days. Inventory is healthy and fully covered.", "🟢 Healthy / Low Risk: ")

    add_sub_header("Multi-Period Horizon Simulation (7d, 14d, 30d)")
    add_body("The engine projects inventory depletion curves across three forward-looking milestones:")
    add_body("Stock(7d) = max(0, Stock - [Daily Velocity × 7] + ReorderLevel × 0.1)\nStock(14d) = max(0, Stock - [14d Demand] + ReorderLevel × 0.1)\nStock(30d) = max(0, Stock - [Daily Velocity × 30] + ReorderLevel × 0.1)", "Simulation Model: ", italic=True)

    # --- Section 7: Operational Guide ---
    add_section_header("7. Step-by-Step Operational Guide for JonBrix Staff")
    add_body("Navigate to the Stock Prediction tab on the sidebar. Review the top summary cards showing total critical items, items needing reorder, and average days to stockout.", "Step 1 — Daily Morning Review: ")
    add_body("Click the 'Critical Items' filter tab to isolate products running out within 7 days.", "Step 2 — Filter High-Risk SKUs: ")
    add_body("Review the suggested 'Recommended Reorder Quantity' column. The numbers are pre-calculated based on 2-week sales demand and momentum.", "Step 3 — Inspect Recommended Quantities: ")
    add_body("Click 'Create Purchase Order' to automatically export the recommended items directly into an approved Purchase Order grouped by supplier.", "Step 4 — 1-Click Purchase Order: ")
    add_body("When the supplier delivers the goods, receive the purchase order in the Inventory module. InvenSight automatically recalculates all velocity curves and clears the alerts.", "Step 5 — Inventory Restock: ")

    # Save Document
    output_path = r"c:\Users\MY PC\Desktop\InvenSight\InvenSight_Stock_Prediction_Guide.docx"
    doc.save(output_path)
    print(f"Successfully generated DOCX at: {output_path}")

if __name__ == "__main__":
    create_document()
